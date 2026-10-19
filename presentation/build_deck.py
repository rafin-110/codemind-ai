"""Build presentation/CodeMind_AI_CA3.pptx (editable) from real screenshots and results.

    pip install python-pptx
    python presentation/build_deck.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "screenshots" / "phase3"
OUT = ROOT / "presentation" / "CodeMind_AI_CA3.pptx"

INK = RGBColor(0x14, 0x1B, 0x2D)
MUTED = RGBColor(0x5B, 0x64, 0x78)
ACCENT = RGBColor(0x4F, 0x46, 0xE5)
GOOD = RGBColor(0x0F, 0x7B, 0x4F)
BAD = RGBColor(0xC2, 0x41, 0x0C)
PANEL = RGBColor(0xF3, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=16, color=INK, bold=False, font="Calibri"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        b = bold
        if isinstance(line, tuple):
            line, b = line
        p.text = line
        p.space_after = Pt(6)
        for r in p.runs:
            r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(size), color, b, font
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def header(slide, kicker, title, n):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    text(slide, 0.6, 0.35, 10, 0.4, kicker.upper(), 12, ACCENT, True)
    text(slide, 0.6, 0.7, 12, 0.9, title, 30, INK, True)
    text(slide, 12.2, 7.0, 0.9, 0.4, f"{n}/8", 11, MUTED)
    text(slide, 0.6, 7.0, 8, 0.4, "CodeMind AI · Agentic Code Review Assistant · CA-3", 11, MUTED)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def card(slide, x, y, w, h, title, body, color=ACCENT, size=13):
    box(slide, x, y, w, h)
    text(slide, x + 0.2, y + 0.12, w - 0.4, 0.4, title, 15, color, True)
    text(slide, x + 0.2, y + 0.55, w - 0.4, h - 0.6, body, size, INK)


# 1 - title + problem + domain
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = INK; bg.line.fill.background()
text(s, 0.8, 0.9, 11, 0.5, "CA-3 · AGENTIC AI · DOMAIN: CODE REVIEW", 14, RGBColor(0xA5, 0xB4, 0xFC), True)
text(s, 0.8, 1.4, 11.5, 1.2, "CodeMind AI", 54, WHITE, True)
text(s, 0.8, 2.5, 11.5, 0.8, "An agentic code review assistant for AI-generated code", 26, RGBColor(0xE0, 0xE7, 0xFF))
text(s, 0.8, 3.7, 7.4, 3, [
    ("The problem", True),
    "Students and “vibe coders” ship AI-generated code they did not write and do not fully understand.",
    "It runs, so it gets accepted, but it hides boundary bugs, swallowed errors and security holes (eval, SQL injection, shell=True, hard-coded keys).",
    "Code review is the normal quality gate, but it does not scale to this volume, and a beginner has no reviewer at all.",
], 16, WHITE)
box(s, 8.7, 3.75, 3.9, 2.6, RGBColor(0x23, 0x2B, 0x45))
text(s, 8.95, 3.9, 3.5, 2.4, [
    ("Why this domain", True),
    "Code is verifiable: it can be parsed, diffed and tested.",
    "So every claim the agents make can be backed by evidence, not just text.",
], 14, WHITE)
notes(s, "30s. Problem: AI writes code people accept without understanding. Code review is the gate, but it doesn't scale "
         "and beginners have no reviewer. Code is a good domain for agents because claims can be checked: parsed, diffed, tested.")

# 2 - challenges vs solution
s = prs.slides.add_slide(BLANK)
header(s, "Existing challenge → proposed solution", "Review like a careful human reviewer, but prove it", 2)
rows = [("Linters give terse warnings", "Each finding has evidence (file + line), why, impact, fix and how it will be validated"),
        ("Chat assistants rewrite whole files", "Minimal single-finding patch shown as a unified diff"),
        ("“Fixed” code is never proven", "Real pytest run on an isolated copy; retry with a different strategy if it fails"),
        ("Tools act without consent", "Hard human approval bound to the exact patch hash; original never written"),
        ("Only one answer offered", "Advisor compares alternatives and trade-offs, incl. “keep current approach”")]
y = 1.75
for a, b in rows:
    box(s, 0.6, y, 4.6, 0.82, RGBColor(0xFD, 0xEE, 0xE6))
    text(s, 0.8, y + 0.2, 4.3, 0.5, a, 15, BAD, True)
    text(s, 5.35, y + 0.17, 0.5, 0.5, "→", 22, MUTED, True)
    box(s, 5.9, y, 6.8, 0.82, RGBColor(0xE7, 0xF5, 0xEE))
    text(s, 6.1, y + 0.2, 6.5, 0.5, b, 15, GOOD, True)
    y += 0.98
text(s, 0.6, 6.7, 12, 0.4, "Core loop:  ANALYZE → EXPLAIN WITH EVIDENCE → COMPARE → HUMAN APPROVAL → APPLY TO COPY → VALIDATE → REPORT",
     14, ACCENT, True)
notes(s, "45s with slide 3. Each existing tool fails in one way; our response to each is a concrete feature you will see in the demo.")

# 3 - agents + orchestration
s = prs.slides.add_slide(BLANK)
header(s, "Agent roles and orchestration", "A planner and five agents share one typed workflow state", 3)
agents = [("1 · Repository Understanding", "README, structure, entry points, dependencies, import graph, doc gaps, with line evidence"),
          ("2 · Bug & Risk Analysis", "Python AST rules R001–R009: division by zero, eval, SQL injection, shell=True, secrets, mutable defaults…"),
          ("3 · Solution & Architecture Advisor", "Current approach vs alternatives and trade-offs; says when current code is fine"),
          ("4 · Code Improvement", "Minimal patch + diff, risks, suggested tests; picks a new strategy from test-failure feedback"),
          ("5 · Critic & Validation", "Baseline tests on original, reviews every patch, then syntax + rule re-check + real pytest on a copy")]
y = 1.7
for t, b in agents:
    card(s, 0.6, y, 6.6, 0.95, t, b, size=12)
    y += 1.05
box(s, 7.6, 1.7, 5.1, 5.15, PANEL)
text(s, 7.85, 1.8, 4.7, 0.4, "Planner / Orchestrator (state machine)", 15, ACCENT, True)
text(s, 7.85, 2.3, 4.7, 4.5, [
    "RECEIVED → PLANNED → ANALYZING → FINDINGS_READY",
    "→ PATCH_PROPOSED → critic review → AWAITING_APPROVAL",
    "→ human approves → APPLYING (copy) → VALIDATING",
    "→ VALIDATED  |  failed → retry new strategy (max 2) → BLOCKED",
    ("Guarantees", True),
    "• only path to applying a patch is human approval",
    "• illegal transitions raise errors (tested)",
    "• every transition logged with its reason",
    "• plan adapts: no tests → baseline skipped; untrusted upload → tests not run",
], 13, INK)
notes(s, "Agentic, not a chatbot: a planner builds a plan from what the repo contains, agents write structured results to "
         "shared state, the critic can send work back, and validation failure feeds the next strategy.")

# 4 - architecture (native shapes so it stays readable and editable; full diagram: docs/architecture.png)
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn


def node(slide, x, y, w, h, label, fill=PANEL, color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, size=11, line=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line or RGBColor(0xC7, 0xCB, 0xD6)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.03))
    for i, part in enumerate(label.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = part; para.alignment = PP_ALIGN.CENTER
        for r in para.runs:
            r.font.size = Pt(size if i == 0 else size - 2); r.font.bold = i == 0
            r.font.color.rgb = color; r.font.name = "Calibri"
    return sh


def arrow(slide, a, b, label=None, color=MUTED, ax="r", bx="l"):
    def pt(sh, side):
        x, y, w, h = sh.left, sh.top, sh.width, sh.height
        return {"r": (x + w, y + h // 2), "l": (x, y + h // 2), "t": (x + w // 2, y), "b": (x + w // 2, y + h)}[side]
    (x1, y1), (x2, y2) = pt(a, ax), pt(b, bx)
    c = slide.shapes.add_connector(1, x1, y1, x2, y2)
    c.line.color.rgb = color; c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln(); tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}); ln.append(tail)
    if label:
        tb = text(slide, 0, 0, 1.6, 0.3, label, 10, color, True)
        tb.left, tb.top = (x1 + x2) // 2 - Inches(0.8), (y1 + y2) // 2 - Inches(0.3)
    return c


s = prs.slides.add_slide(BLANK)
header(s, "Complete architecture", "From upload to verified report", 4)
LIGHT_A = RGBColor(0xEE, 0xF0, 0xFF); LIGHT_G = RGBColor(0xE7, 0xF5, 0xEE); LIGHT_Y = RGBColor(0xFF, 0xF7, 0xE0)
dev = node(s, 0.5, 1.75, 1.5, 0.75, "Developer\n(human reviewer)", LIGHT_Y)
ui_ = node(s, 2.45, 1.75, 1.9, 0.75, "Streamlit UI\n7 views")
ing = node(s, 4.8, 1.75, 3.6, 0.75, "Repository ingestion\nunsafe ZIP (traversal, symlink, bomb) → BLOCKED")
orig = node(s, 8.85, 1.75, 3.95, 0.75, "Original repo\nread-only · hashed · never written", LIGHT_G)
arrow(s, dev, ui_); arrow(s, ui_, ing); arrow(s, ing, orig)
plan = node(s, 0.5, 2.95, 7.0, 0.65, "Planner / Orchestrator · state machine · decision log", LIGHT_A, ACCENT, size=12, line=ACCENT)
state = node(s, 7.85, 2.95, 4.95, 0.65, "Shared structured findings & workflow state", LIGHT_A, ACCENT, size=12, line=ACCENT)
arrow(s, plan, state)
a1 = node(s, 0.5, 4.0, 1.9, 0.85, "A1 Repository\nUnderstanding")
a2 = node(s, 2.6, 4.0, 1.9, 0.85, "A2 Bug & Risk\nAST rules R001–R009")
a5b = node(s, 4.7, 4.0, 1.9, 0.85, "A5 Baseline tests\nconfirm findings")
a3 = node(s, 6.8, 4.0, 1.9, 0.85, "A3 Solution &\nArchitecture Advisor")
a4 = node(s, 8.9, 4.0, 1.9, 0.85, "A4 Code\nImprovement (patch)")
a5r = node(s, 11.0, 4.0, 1.8, 0.85, "A5 Critic\nscope · compile")
for x, y in ((a1, a2), (a2, a5b), (a5b, a3), (a3, a4), (a4, a5r)):
    arrow(s, x, y)
tools = node(s, 0.5, 5.25, 2.9, 0.75, "Tools: ast · difflib · compile()\npytest subprocess (timeout, env)", RGBColor(0xF5, 0xF5, 0xF5))
rep = node(s, 3.7, 5.25, 1.6, 0.75, "Final report\nJSON · Markdown", LIGHT_G)
val = node(s, 5.8, 5.25, 2.0, 0.75, "A5 Validate\nsyntax · re-check · tests")
copy = node(s, 8.3, 5.25, 2.0, 0.75, "Isolated working copy\npatch applied here only", LIGHT_G)
appr = node(s, 10.6, 5.1, 2.2, 1.05, "Human approval\nbound to patch hash", LIGHT_Y, shape=MSO_SHAPE.DIAMOND, size=10)
arrow(s, a5r, appr, "accepted", ax="b", bx="t")
arrow(s, appr, copy, "approve", GOOD, ax="l", bx="r"); arrow(s, copy, val, ax="l", bx="r")
arrow(s, val, rep, "passed", GOOD, ax="l", bx="r")
arrow(s, val, a4, "failed → retry (max 2)", BAD, ax="t", bx="b")
back = node(s, 10.7, 6.4, 2.0, 0.5, "Reject → back to findings\noriginal untouched", RGBColor(0xFD, 0xEE, 0xE6), BAD, size=10)
arrow(s, appr, back, None, BAD, ax="b", bx="t")
blk = node(s, 5.85, 6.4, 1.9, 0.5, "BLOCKED → escalate\nretries exhausted", RGBColor(0xFD, 0xEE, 0xE6), BAD, size=10)
arrow(s, val, blk, None, BAD, ax="b", bx="t")
notes(s, "Walk top to bottom: upload is checked, the original is frozen, the planner runs the agents over shared state. "
         "The only way to a changed file is the approval diamond, and changes go to a separate copy. Failed validation "
         "loops back to the improvement agent; after 2 retries it is BLOCKED. Full Mermaid diagram: docs/architecture.png.")

# 5 - demo
s = prs.slides.add_slide(BLANK)
header(s, "Working Streamlit PoC (synthetic data)", "Demo: confirmed bug → failed fix → retry → verified", 5)
s.shapes.add_picture(str(SHOTS / "06_findings_F004_confirmed.png"), Inches(0.6), Inches(1.7), width=Inches(6.0))
s.shapes.add_picture(str(SHOTS / "11_F004_attempt2_after_failed_validation.png"), Inches(6.75), Inches(1.7), width=Inches(6.0))
text(s, 0.6, 5.55, 6.0, 1.3, ["F-004 division by zero is Confirmed: the sample's own test fails on that exact line in the baseline run."], 13, MUTED)
text(s, 6.75, 5.55, 6.0, 1.3, ["Attempt 1 (return 0.0) failed a real test, so the agent proposes “Raise ValueError”. "
                               "Approved → 9/9 tests pass, original unchanged."], 13, MUTED)
notes(s, "2 minutes LIVE DEMO (see DEMO_SCRIPT.md). Backup: these are screenshots from a real run.")

# 6 - validation & results
s = prs.slides.add_slide(BLANK)
header(s, "Output validation and actual results", "Every number below is from a real run", 6)
stats = [("105", "automated tests passed", "Windows 11, Python 3.13"), ("9 / 9", "planted issues found", "exact rule+file+line, 0 extra"),
         ("8 / 9", "fixes validated", "7 first try, 1 after retry"), ("0", "originals modified", "re-hashed after every run")]
x = 0.6
for big, lab, sub in stats:
    box(s, x, 1.75, 2.95, 1.75)
    text(s, x + 0.2, 1.85, 2.6, 0.8, big, 36, ACCENT, True)
    text(s, x + 0.2, 2.7, 2.6, 0.4, lab, 14, INK, True)
    text(s, x + 0.2, 3.05, 2.6, 0.4, sub, 11, MUTED)
    x += 3.08
text(s, 0.6, 3.8, 6.2, 3, [
    ("How results are validated", True),
    "• statuses: Passed / Failed / Not run / Blocked / Partially verified",
    "• pass = compiles + rule no longer matches + no regressions vs baseline + linked tests pass",
    "• browser end-to-end run in real Chrome with assertions on the downloaded report",
    "• fresh install from requirements.txt: 104 passed, 1 optional skip",
], 13)
text(s, 7.0, 3.8, 5.7, 3, [
    ("Found and fixed in testing", True),
    "First Windows run: 4 tests failed. Pytest prints grades\\stats.py on Windows, so the bug was not linked to its failing test. "
    "Fixed (one line) + regression test. Linux runs could not have caught it.",
    ("Honest caveat", True),
    "9/9 is on our synthetic sample, so it shows the rules work as designed, not real-world accuracy.",
], 13)
notes(s, "45s. Emphasise: no number is estimated; evidence files are in docs/test-evidence. Mention the Windows bug as proof the testing was real.")

# 7 - limitations & future
s = prs.slides.add_slide(BLANK)
header(s, "Limitations and future scope", "What it cannot do yet, and where it goes next", 7)
card(s, 0.6, 1.75, 5.9, 5.0, "Limitations", [
    "• 9 deterministic rules, Python only; no cross-function data flow",
    "• fix templates for 8 rules; R008 has no auto-fix by design",
    "• explanations from a knowledge base (labelled Mock LLM); Ollama optional, only rewords",
    "• tests run in a subprocess, not a security sandbox; Windows enforces only a timeout",
    "• validation is only as good as the repo's own tests",
    "• single user, no database, no GitHub URL input",
], BAD, 14)
card(s, 6.8, 1.75, 5.9, 5.0, "Future scope", [
    "• LLM reviewer for fixes beyond templates, kept behind the same critic + approval + test gates",
    "• Pull-request integration: comment on changed lines, suggest patches, never push",
    "• Container / microVM sandbox for safe testing of any upload",
    "• Bandit, Semgrep, Ruff, ESLint as tools for the risk agent; more languages",
    "• Propose a failing test first when the repo has none",
    "• Evaluate on real labelled bug datasets",
], GOOD, 14)
notes(s, "45s. Say the limitations plainly; the design keeps the guardrails when an LLM is added later.")

# 8 - close
s = prs.slides.add_slide(BLANK)
header(s, "Summary", "Understand → evidence → compare → approve → verify", 8)
text(s, 0.6, 1.9, 12, 4, [
    "• Agentic code review with a planner, five agents, a critic loop and bounded retries",
    "• Every finding cites real lines; confirmed findings are reproduced by a real test",
    "• Nothing changes without the developer's approval of the exact patch, and the original is never touched",
    "• Fixes are proven by real test runs, and failures are reported, not hidden",
    "• Fully offline, no paid API: Python, Streamlit, ast, pytest",
], 20)
text(s, 0.6, 5.6, 12, 0.8, "Thank you · Questions?", 30, ACCENT, True)
notes(s, "Close in 15s and invite questions. Viva prep: presentation/VIVA_QA.md.")

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(f"wrote {OUT} ({len(prs.slides)} slides)")
